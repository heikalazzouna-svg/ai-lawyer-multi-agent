"""
AI Lawyer API Bridge — Powered by LangGraph Multi-Agent System.

This FastAPI server connects the Next.js React frontend to the
LangGraph multi-agent pipeline. Each user query flows through:

    Supervisor → Specialist Agent → Citator → Response

Endpoints:
    POST /upload       — Upload PDFs, index them, run proactive scan
    POST /chat         — Run the full multi-agent pipeline on a query (blocking)
    POST /chat/stream  — Run the pipeline with real-time SSE streaming
    GET  /pdfs/*       — Serve uploaded PDFs for the React PDF viewer
"""

import os
import shutil
import uuid
import smtplib
from email.message import EmailMessage
from typing import List, Optional
from fastapi import FastAPI, UploadFile, File, HTTPException, Depends, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel
from sqlalchemy.orm import Session
import json
import io
import docx
from datetime import timedelta

from vector_database import index_pdf, retrieve_docs, PDFS_DIRECTORY
from rag_pipeline import proactive_scan
from agents.graph import create_legal_agent_graph
from agents.stream_graph import stream_legal_pipeline
from database import Base, engine, get_db
from models import ChatSession, ChatMessage, User
from auth import (
    get_password_hash, 
    verify_password, 
    create_access_token, 
    ACCESS_TOKEN_EXPIRE_MINUTES, 
    get_current_user_email
)

app = FastAPI(
    title="AI Lawyer — Multi-Agent API",
    description="LangGraph-powered legal document analysis with Supervisor routing.",
    version="2.0.0",
)

# Enable CORS for Next.js
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Ensure PDFs directory exists
os.makedirs(PDFS_DIRECTORY, exist_ok=True)

# Mount the PDFs directory so Next.js can embed the PDFs in an iframe
app.mount("/pdfs", StaticFiles(directory=PDFS_DIRECTORY), name="pdfs")

# Compile the LangGraph once at startup
legal_graph = create_legal_agent_graph()

# Initialize the Database Tables
Base.metadata.create_all(bind=engine)


class ChatRequest(BaseModel):
    query: str
    filenames: List[str]
    session_id: Optional[int] = None

class ExportRequest(BaseModel):
    content: str
    title: str = "Legal_Draft"

class UserCreate(BaseModel):
    first_name: str
    last_name: str
    email: str
    password: str

class Token(BaseModel):
    access_token: str
    token_type: str


@app.post("/register")
async def register(user: UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.email == user.email).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    hashed_password = get_password_hash(user.password)
    verification_token = str(uuid.uuid4())
    
    new_user = User(
        email=user.email,
        first_name=user.first_name,
        last_name=user.last_name,
        hashed_password=hashed_password,
        is_verified=False,
        verification_token=verification_token
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    
    # ─── REAL EMAIL SENDING ───
    verify_link = f"http://localhost:3000/?token={verification_token}"
    
    sender_email = os.getenv("SMTP_EMAIL")
    sender_password = os.getenv("SMTP_PASSWORD")
    
    if sender_email and sender_password:
        try:
            msg = EmailMessage()
            msg.set_content(f"Hello {user.first_name},\n\nPlease verify your AI Lawyer account by clicking this link:\n{verify_link}\n\nThank you,\nAI Lawyer Team")
            msg['Subject'] = 'Verify your AI Lawyer Account'
            msg['From'] = sender_email
            msg['To'] = user.email

            server = smtplib.SMTP("smtp.gmail.com", 587)
            server.starttls()
            server.login(sender_email, sender_password)
            server.send_message(msg)
            server.quit()
            print(f"📧 REAL EMAIL SENT TO: {user.email}")
        except Exception as e:
            print(f"❌ Failed to send email: {e}")
            raise HTTPException(status_code=500, detail=f"Failed to send email: {str(e)}")
    else:
        # Fallback if they didn't set up credentials yet
        print("\n" + "="*50)
        print("📧 [MOCK] EMAIL SENT TO:", user.email)
        print("Subject: Verify your AI Lawyer Account")
        print(f"Body: Hello {user.first_name},\n\nPlease verify your account by clicking this link:\n{verify_link}")
        print("="*50 + "\n")
        print("WARNING: You must set SMTP_EMAIL and SMTP_PASSWORD in .env to send real emails.")
    
    return {"message": "Account created! Please check your email to verify your account."}

@app.get("/verify/{token}")
async def verify_email(token: str, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.verification_token == token).first()
    if not user:
        raise HTTPException(status_code=400, detail="Invalid verification token.")
        
    user.is_verified = True
    user.verification_token = None
    db.commit()
    
    return {"message": "Email successfully verified. You may now log in."}


@app.post("/token", response_model=Token)
async def login(form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == form_data.username).first()
    if not user or not verify_password(form_data.password, user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect email or password",
            headers={"WWW-Authenticate": "Bearer"},
        )
        
    if not user.is_verified:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Please check your email and verify your account before logging in.",
        )
    access_token_expires = timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = create_access_token(
        data={"sub": user.email}, expires_delta=access_token_expires
    )
    return {"access_token": access_token, "token_type": "bearer"}


@app.post("/upload")
async def upload_documents(files: List[UploadFile] = File(...), email: str = Depends(get_current_user_email)):
    """Upload PDFs, index them into the vector store, and run a proactive legal scan."""
    uploaded_names = []
    
    for file in files:
        if not file.filename.endswith('.pdf'):
            continue
            
        file_path = os.path.join(PDFS_DIRECTORY, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        index_pdf(file_path)
        uploaded_names.append(file.filename)
        
    if not uploaded_names:
        raise HTTPException(status_code=400, detail="No valid PDFs uploaded.")
        
    # Run proactive scan
    scan_result_text = "No critical clauses found."
    try:
        scan_docs = retrieve_docs(
            "Governing law, jurisdiction, term, termination, liability, obligations",
            uploaded_names
        )
        if scan_docs:
            scan_result = proactive_scan(scan_docs)
            scan_result_text = str(scan_result.content) if hasattr(scan_result, 'content') else str(scan_result)
    except Exception as e:
        scan_result_text = f"Scan failed: {str(e)}"
        
    return {
        "message": f"Successfully indexed {len(uploaded_names)} documents.",
        "filenames": uploaded_names,
        "scan_result": scan_result_text,
    }


@app.post("/chat")
async def chat_endpoint(request: ChatRequest):
    """
    Run the full LangGraph multi-agent pipeline.
    
    Flow: Supervisor → [Researcher | Risk Assessor | Drafter] → Citator → Response
    """
    if not request.filenames:
        raise HTTPException(status_code=400, detail="No documents selected.")
    
    try:
        # Retrieve relevant documents from the vector store
        retrieved_docs = retrieve_docs(request.query, request.filenames)
        
        # Build the initial state for the LangGraph
        initial_state = {
            "query": request.query,
            "documents": retrieved_docs,
            "filenames": request.filenames,
            "context": "",
            "next_agent": "",
            "agent_output": "",
            "final_answer": "",
            "agent_name": "",
            "citations_verified": False,
        }
        
        # Run the full multi-agent graph
        result = legal_graph.invoke(initial_state)
        
        # Extract the final answer (Citator output takes priority)
        answer = result.get("final_answer") or result.get("agent_output", "No response generated.")
        agent_name = result.get("agent_name", "Unknown")
        citations_verified = result.get("citations_verified", False)
        
        return {
            "answer": answer,
            "agent": agent_name,
            "citations_verified": citations_verified,
        }
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/chat/stream")
async def chat_stream_endpoint(
    request: ChatRequest, 
    db: Session = Depends(get_db), 
    email: str = Depends(get_current_user_email)
):
    """
    Run the full multi-agent pipeline with real-time Server-Sent Events.

    Emits events:
        trace — Agent status updates (supervisor routing, specialist working, citator verifying)
        token — Individual text chunks streamed from the specialist LLM
        done  — Final result with agent name and citation verification status
    """
    if not request.filenames:
        raise HTTPException(status_code=400, detail="No documents selected.")

    try:
        retrieved_docs = retrieve_docs(request.query, request.filenames)

        initial_state = {
            "query": request.query,
            "documents": retrieved_docs,
            "filenames": request.filenames,
            "context": "",
            "next_agent": "",
            "agent_output": "",
            "final_answer": "",
            "agent_name": "",
            "citations_verified": False,
        }

        # 1. Fetch or create a DB session for this chat interaction linked to the user
        chat_session = None
        if request.session_id:
            chat_session = db.query(ChatSession).filter(
                ChatSession.id == request.session_id, 
                ChatSession.user_id == email
            ).first()
        
        if not chat_session:
            chat_session = ChatSession(user_id=email)
            db.add(chat_session)
            db.commit()
            db.refresh(chat_session)

        # 2. Save the User's query immediately
        user_msg = ChatMessage(
            session_id=chat_session.id,
            role="user",
            content=request.query
        )
        db.add(user_msg)
        db.commit()

        # 3. Create a wrapper generator to intercept the final answer
        async def stream_and_save():
            final_answer = "No response generated."
            agent_name = "Unknown"
            citations_verified = False

            async for event_str in stream_legal_pipeline(initial_state):
                
                # Parse the event to catch the "done" payload
                if event_str.startswith("event: done"):
                    try:
                        # Extract the JSON data part (skip 'event: done\ndata: ')
                        data_part = event_str.split("data: ")[1].strip()
                        payload = json.loads(data_part)
                        
                        final_answer = payload.get("final_answer", final_answer)
                        agent_name = payload.get("agent", agent_name)
                        citations_verified = payload.get("citations_verified", citations_verified)
                        
                        # Inject session_id and yield the modified event
                        payload["session_id"] = chat_session.id
                        yield f"event: done\ndata: {json.dumps(payload)}\n\n"
                    except Exception:
                        yield event_str
                else:
                    yield event_str

            # Once the stream is finished, save the AI's final response to the database
            ai_msg = ChatMessage(
                session_id=chat_session.id,
                role="ai",
                content=final_answer,
                agent_name=agent_name,
                citations_verified=citations_verified
            )
            db.add(ai_msg)
            db.commit()

        return StreamingResponse(
            stream_and_save(),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            },
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/history")
async def get_chat_history(db: Session = Depends(get_db), email: str = Depends(get_current_user_email)):
    """Retrieve all past chat sessions and their messages, scoped to the user."""
    sessions = db.query(ChatSession).filter(ChatSession.user_id == email).order_by(ChatSession.created_at.desc()).all()
    
    result = []
    for s in sessions:
        msgs = db.query(ChatMessage).filter(ChatMessage.session_id == s.id).order_by(ChatMessage.created_at.asc()).all()
        result.append({
            "session_id": s.id,
            "created_at": s.created_at,
            "messages": [
                {
                    "role": m.role,
                    "content": m.content,
                    "agent": m.agent_name,
                    "verified": m.citations_verified
                } for m in msgs
            ]
        })
    return {"history": result}


@app.post("/export/docx")
async def export_to_docx(request: ExportRequest, email: str = Depends(get_current_user_email)):
    """Generate a Microsoft Word (.docx) file from the provided markdown/text."""
    try:
        doc = docx.Document()
        doc.add_heading(request.title, 0)
        
        # Super simple markdown-to-docx parser (just handling paragraphs)
        paragraphs = request.content.split("\n\n")
        for p in paragraphs:
            if p.strip():
                doc.add_paragraph(p.strip())
                
        # Save to BytesIO stream
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{request.title}.docx"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate docx: {str(e)}")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
